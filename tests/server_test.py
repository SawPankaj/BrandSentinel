import io
import tempfile
import unittest
import sys
from unittest import mock
from pathlib import Path
from email.message import Message

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
import server

class ServerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=server.ROOT)
        server.DATA_DIR = Path(self.temp.name)
        server.DB_PATH = Path(self.temp.name) / "test.sqlite3"

    def tearDown(self): self.temp.cleanup()

    def test_state_is_partitioned_by_workspace(self):
        server.save_state("alpha", {"rules": [{"id":"A","title": "A"}], "reviews": [], "audit": []})
        server.save_state("beta", {"rules": [{"id":"B","title": "B"}], "reviews": [], "audit": []})
        self.assertEqual(server.load_state("alpha")["rules"][0]["title"], "A")
        self.assertEqual(server.load_state("beta")["rules"][0]["title"], "B")

    def test_workspace_id_rejects_path_like_values(self):
        with self.assertRaises(ValueError): server.valid_workspace("../outside")

    def test_audit_is_derived_server_side_and_append_only(self):
        first={"rules":[],"audit":[],"reviews":[{"id":"R1","title":"First","text":"one","status":"In review"},{"id":"R2","title":"Second","text":"two","status":"In review"}]}
        server.save_state("alpha",first,"A User","Brand manager")
        first["audit"]=[{"text":"forged"}]
        first["reviews"].append({"id":"R3","title":"Third","text":"three","status":"In review"})
        server.save_state("alpha",first,"B User","Brand manager")
        self.assertEqual(len(server.audit_entries("alpha")),3)
        self.assertFalse(any("forged" in row["action"] for row in server.audit_entries("alpha")))

    def test_review_versions_capture_content_changes(self):
        state = {"rules": [], "audit": [], "reviews": [{"id": "R1", "title": "Draft", "text": "one"}]}
        server.save_state("alpha", state)
        state["reviews"][0]["text"] = "two"
        server.save_state("alpha", state)
        versions = server.review_versions("alpha", "R1")
        self.assertEqual([x["version"] for x in versions], [2, 1])
        self.assertNotEqual(versions[0]["content_hash"], versions[1]["content_hash"])

    def test_server_rejects_approval_with_unresolved_findings(self):
        state = {"rules": [], "audit": [], "reviews": [{"id": "R2", "title": "Claims", "status": "Approved", "findings": [{"id":"F1","status": "Escalated", "source": "Claims register"}]}]}
        with self.assertRaises(ValueError): server.save_state("alpha", state)

    def test_server_requires_legal_clearance_for_high_risk_findings(self):
        state = {"rules": [], "audit": [], "reviews": [{"id": "R4", "title": "Claim", "status": "Approved", "findings": [{"id":"F1","kind": "Unsupported claim", "status": "Accepted", "source": "Claims register"}]}]}
        with self.assertRaises(ValueError): server.save_state("alpha", state)

    def test_server_locks_approved_content(self):
        state = {"rules": [], "audit": [], "reviews": [{"id": "R3", "title": "Clean", "status": "In review", "text": "Initial"}]}
        server.save_state("alpha", state)
        state["reviews"][0]["status"] = "Approved"
        server.save_state("alpha", state)
        state["reviews"][0]["text"] = "Changed after approval"
        with self.assertRaises(server.ServiceError): server.save_state("alpha", state)

    def test_server_locks_approved_metadata_and_findings(self):
        state={"rules":[],"audit":[],"reviews":[{"id":"R5","title":"Locked title","status":"In review","text":"Safe copy","findings":[]}]}
        server.save_state("alpha",state)
        state["reviews"][0]["status"]="Approved"
        server.save_state("alpha",state)
        state["reviews"][0]["title"]="Tampered title"
        with self.assertRaises(server.ServiceError): server.save_state("alpha",state)

    def test_workspace_switch_rotates_csrf_and_stays_with_membership(self):
        server.create_initial_admin("admin@example.test","Admin","correct horse battery","Main")
        owner=server.login_user("admin@example.test","correct horse battery")
        current=server.resolve_session(owner[1])
        old_csrf=current["csrf_token"]
        extra=server.create_workspace_for_user(current["user_id"],"Second workspace")
        new_csrf=server.switch_session_workspace(owner[1],extra["id"])
        updated=server.resolve_session(owner[1])
        self.assertEqual(updated["workspace_id"],extra["id"])
        self.assertNotEqual(old_csrf,new_csrf)
        self.assertEqual(updated["csrf_token"],new_csrf)

    def test_password_hashing_and_production_key_requirement(self):
        self.assertTrue(server.security.verify_password("correct horse battery",server.security.hash_password("correct horse battery")))
        with mock.patch.dict("os.environ",{"BRANDSENTINEL_MASTER_KEY":""}):
            with self.assertRaises(RuntimeError): server.security.load_master_key(Path(self.temp.name)/"prod",production=True)

    def test_hostile_ids_are_rejected_before_rendering(self):
        payload={"rules":[],"reviews":[{"id":"bad-id-<script>","title":"Bad","findings":[]}],"audit":[]}
        with self.assertRaises(ValueError): server.normalize_state(payload)

    def test_source_login_throttle_cannot_be_evaded_by_rotating_emails(self):
        with mock.patch.object(server.security,"verify_password",return_value=False):
            for number in range(25):
                with self.assertRaises(server.ServiceError): server.login_user(f"person{number}@example.test","wrong password","192.0.2.10")
            with self.assertRaises(server.ServiceError) as exc: server.login_user("other@example.test","wrong password","192.0.2.10")
            self.assertEqual(exc.exception.status,429)

    def test_text_and_html_extraction_excludes_scripts(self):
        self.assertEqual(server.extract_text("draft.txt", b"hello\nworld")["text"], "hello\nworld")
        parsed = server.extract_text("draft.html", b"<h1>Welcome</h1><script>steal()</script><p>Students</p>")["text"]
        self.assertIn("Welcome", parsed); self.assertIn("Students", parsed); self.assertNotIn("steal", parsed)

    def test_html_extract_reports_language_alt_link_and_heading_issues(self):
        parsed=server.extract_text("page.html",b"<html><h1>Welcome</h1><h3>Details</h3><a href='/x'>Click here</a><img src='hero.png'></html>")
        reasons=[item["reason"] for item in parsed["checks"]]
        self.assertEqual(len(parsed["checks"]),4)
        self.assertTrue(any("language" in reason for reason in reasons))
        self.assertTrue(any("alternative text" in reason for reason in reasons))
        self.assertTrue(any("link" in reason for reason in reasons))
        self.assertTrue(any("Heading levels" in reason for reason in reasons))

    def test_multipart_upload_parser_keeps_filename_and_bytes(self):
        boundary = "----brandsentinel-test"
        body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"draft.txt\"\r\nContent-Type: text/plain\r\n\r\nSubmitted words\r\n--{boundary}--\r\n").encode()
        filename, data = server.parse_multipart(f"multipart/form-data; boundary={boundary}", body)
        self.assertEqual(filename, "draft.txt")
        self.assertEqual(data, b"Submitted words")

    def test_secure_handler_denies_anonymous_state_access(self):
        class Harness(server.SecureHandler):
            def __init__(self,method,path,body=b""):
                self.command=method; self.path=path; self.headers=Message(); self.headers["Content-Length"]=str(len(body)); self.rfile=io.BytesIO(body); self.wfile=io.BytesIO(); self.status=None
            def cookie_token(self): return None
            def send_response(self,status,message=None): self.status=status
            def send_header(self,*_): pass
            def end_headers(self): pass
        request=Harness("GET","/api/state")
        request.do_GET()
        self.assertEqual(request.status,401)

    def test_secure_handler_round_trips_authenticated_state(self):
        class Harness(server.SecureHandler):
            def __init__(self,method,path,body=b""):
                self.command=method; self.path=path; self.headers=Message(); self.headers["Content-Length"]=str(len(body)); self.headers["Content-Type"]="application/json"; self.rfile=io.BytesIO(body); self.wfile=io.BytesIO(); self.status=None
            def require_session(self): return {"workspace_id":"http-test","name":"Test User","role":"Brand manager"}
            def require_csrf(self,current): pass
            def verify_origin(self): pass
            def send_response(self,status,message=None): self.status=status
            def send_header(self,*_): pass
            def end_headers(self): pass
        payload={"rules":[{"id":"K1","title":"HTTP rule"}],"reviews":[],"audit":[],"_revision":""}
        put=Harness("PUT","/api/state",server.json.dumps(payload).encode()); put.do_PUT()
        self.assertEqual(put.status,200)
        get=Harness("GET","/api/state"); get.do_GET()
        self.assertEqual(get.status,200)
        self.assertEqual(server.json.loads(get.wfile.getvalue())["rules"][0]["title"],"HTTP rule")

    def test_static_server_does_not_expose_database_or_source_files(self):
        class Harness(server.SecureHandler):
            def __init__(self,path):
                self.command="GET"; self.path=path; self.headers=Message(); self.wfile=io.BytesIO(); self.status=None
            def send_response(self,status,message=None): self.status=status
            def send_header(self,*_): pass
            def end_headers(self): pass
        for path in ("/server.py", "/backend/server.py", "/data/brandsentinel.sqlite3", "/tests/server_test.py"):
            request=Harness(path); request.do_GET(); self.assertEqual(request.status,404)

    def test_static_server_serves_moved_frontend_assets(self):
        class Harness(server.SecureHandler):
            def __init__(self,path):
                self.command="GET"; self.path=path; self.headers=Message(); self.wfile=io.BytesIO(); self.status=None
            def send_response(self,status,message=None): self.status=status
            def send_header(self,*_): pass
            def end_headers(self): pass
        for path,content in (("/","id=\"login-form\""),("/styles.css",".auth-screen"),("/app.js","function boot()"),("/core.js","runChecks")):
            request=Harness(path); request.do_GET()
            self.assertEqual(request.status,200,path)
            self.assertIn(content,request.wfile.getvalue().decode("utf-8"),path)

    def test_workspace_state_is_encrypted_at_rest(self):
        state={"rules":[],"reviews":[{"id":"PRIVATE","title":"Sensitive title","text":"Sensitive draft text"}],"audit":[]}
        server.save_state("tenant-one",state)
        with server.db_session() as db:
            stored=db.execute("SELECT state_json FROM workspace_state WHERE workspace_id='tenant-one'").fetchone()[0]
        self.assertNotIn("Sensitive draft text",stored)
        self.assertEqual(server.load_state("tenant-one")["reviews"][0]["text"],"Sensitive draft text")
        with self.assertRaises(Exception): server.security.decrypt_text(server.get_master_key(),"tenant-two","state",stored)

    def test_stale_revision_cannot_overwrite_a_concurrent_change(self):
        server.save_state("tenant",{"rules":[],"reviews":[],"audit":[]})
        stale=server.load_state("tenant")
        server.save_state("tenant",{"rules":[{"id":"K1","title":"Current"}],"reviews":[],"audit":[]})
        stale["rules"].append({"id":"K2","title":"Stale write"})
        with self.assertRaises(server.ServiceError) as exc: server.save_state("tenant",stale)
        self.assertEqual(exc.exception.status,409)
        self.assertEqual([r["id"] for r in server.load_state("tenant")["rules"]],["K1"])

    def test_creator_cannot_add_brand_knowledge(self):
        payload={"rules":[{"id":"K1","title":"Unauthorized change"}],"reviews":[],"audit":[]}
        with self.assertRaises(server.ServiceError) as exc: server.save_state("tenant",payload,"Creator","Content creator")
        self.assertEqual(exc.exception.status,403)

    def test_server_recreates_missing_checks_and_reopens_resolved_findings_on_new_text(self):
        state={"rules":[],"reviews":[{"id":"R1","title":"Draft","text":"We guarantee 100% placement.","status":"Needs changes","findings":[]}],"audit":[]}
        server.save_state("tenant",state,"Creator","Content creator")
        saved=server.load_state("tenant")
        self.assertTrue(any(f["kind"]=="Unsupported claim" for f in saved["reviews"][0]["findings"]))
        claim=next(f for f in saved["reviews"][0]["findings"] if f["kind"]=="Unsupported claim")
        self.assertEqual(claim["source"],"Claims library · No approved claim found")
        claim["status"]="Accepted"; claim["note"]="Checked by creator"
        server.save_state("tenant",saved,"Creator","Content creator")
        updated=server.load_state("tenant")
        updated["reviews"][0]["text"]="We guarantee 100% placement again."
        server.save_state("tenant",updated,"Creator","Content creator")
        self.assertTrue(any(f["kind"]=="Unsupported claim" and f["status"]=="Open" for f in server.load_state("tenant")["reviews"][0]["findings"]))

    def test_decision_rationale_cannot_be_changed_without_a_new_transition(self):
        state={"rules":[],"reviews":[{"id":"R6","title":"Copy","text":"Safe","status":"In review","findings":[{"id":"F1","kind":"Tone of voice","source":"Voice guide","status":"Open"}]}],"audit":[]}
        server.save_state("tenant",state,"Creator","Content creator")
        accepted=server.load_state("tenant")
        accepted["reviews"][0]["findings"][0].update(status="Dismissed",note="Not applicable")
        server.save_state("tenant",accepted,"Creator","Content creator")
        tampered=server.load_state("tenant")
        tampered["reviews"][0]["findings"][0]["note"]="Evidence removed"
        with self.assertRaises(server.ServiceError): server.save_state("tenant",tampered,"Creator","Content creator")

    def test_team_member_role_updates_preserve_last_administrator(self):
        server.create_initial_admin("owner@example.test","Owner","correct horse battery","Test workspace")
        workspace=server.membership_list(1)[0]["id"]
        initial=server.workspace_users(workspace)[0]
        with self.assertRaises(server.ServiceError) as exc: server.update_member_role(workspace,initial["id"],"Content creator")
        self.assertEqual(exc.exception.status,409)
        member=server.create_member(workspace,"member@example.test","Member User","long first password","Content creator")
        updated=server.update_member_role(workspace,member["id"],"Legal reviewer")
        self.assertEqual(updated["role"],"Legal reviewer")

    def test_csrf_token_is_required_for_authenticated_mutations(self):
        class Harness(server.SecureHandler):
            def verify_origin(self): pass
        request=Harness.__new__(Harness)
        request.headers=Message()
        with self.assertRaises(server.ServiceError) as exc: request.require_csrf({"csrf_token":"expected-token"})
        self.assertEqual(exc.exception.status,403)

    def test_failed_login_lockout_is_persisted(self):
        server.create_initial_admin("owner@example.test","Owner User","correct horse battery","Test workspace")
        with mock.patch.object(server.security,"verify_password",return_value=False):
            for _ in range(5):
                with self.assertRaises(server.ServiceError) as exc: server.login_user("owner@example.test","wrong password")
                self.assertEqual(exc.exception.status,401)
            with self.assertRaises(server.ServiceError) as exc: server.login_user("owner@example.test","wrong password")
            self.assertEqual(exc.exception.status,429)

    def test_word_powerpoint_and_pdf_extraction(self):
        from docx import Document
        from pptx import Presentation
        from reportlab.pdfgen import canvas
        word=io.BytesIO(); doc=Document(); doc.add_paragraph("Word body text"); doc.save(word)
        slides=io.BytesIO(); deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[5]); slide.shapes.title.text="Slide headline"; deck.save(slides)
        pdf=io.BytesIO(); page=canvas.Canvas(pdf); page.drawString(72,720,"PDF body text"); page.save()
        self.assertIn("Word body text",server.extract_text("draft.docx",word.getvalue())["text"])
        self.assertIn("Slide headline",server.extract_text("draft.pptx",slides.getvalue())["text"])
        self.assertIn("PDF body text",server.extract_text("draft.pdf",pdf.getvalue())["text"])

    def test_unsupported_and_invalid_files_fail(self):
        with self.assertRaises(ValueError): server.extract_text("image.png",b"x")
        with self.assertRaises(Exception): server.extract_text("scan.pdf",b"not a PDF")

if __name__ == "__main__": unittest.main(verbosity=2)
