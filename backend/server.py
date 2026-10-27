"""BrandSentinel application server with authenticated, workspace-scoped access."""
from __future__ import annotations

from collections import Counter
import io
import json
import re
import sqlite3
import threading
import zipfile
import os
import hashlib
import ipaddress
import secrets
import getpass
import argparse
from http.cookies import SimpleCookie
from datetime import timedelta
from contextlib import contextmanager
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
import security

ROOT = Path(__file__).resolve().parent
FRONTEND_DIR = ROOT.parent / "frontend"
DATA_DIR = ROOT / "data"
DB_PATH = DATA_DIR / "brandsentinel.sqlite3"
DEFAULT_WORKSPACE = "northstar-demo"
MAX_BODY = 16 * 1024 * 1024
ALLOWED = {".txt", ".md", ".html", ".htm", ".pdf", ".docx", ".pptx"}
STATIC_FILES = {"index.html", "styles.css", "app.js", "core.js"}
LOCK = threading.RLock()
MASTER_KEY: bytes | None = None
SESSION_HOURS = 8
COOKIE_NAME = "brandsentinel_session"


def get_master_key() -> bytes:
    global MASTER_KEY
    if MASTER_KEY is None:
        MASTER_KEY = security.load_master_key(DATA_DIR, production=os.environ.get("BRANDSENTINEL_ENV") == "production")
    return MASTER_KEY


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def connect() -> sqlite3.Connection:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB_PATH, timeout=10)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    db.execute("PRAGMA journal_mode=WAL")
    db.executescript("""
      CREATE TABLE IF NOT EXISTS workspace_state (
        workspace_id TEXT PRIMARY KEY,
        state_json TEXT NOT NULL,
        updated_at TEXT NOT NULL
      );
      CREATE TABLE IF NOT EXISTS audit_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        workspace_id TEXT NOT NULL,
        created_at TEXT NOT NULL,
        actor TEXT NOT NULL,
        action TEXT NOT NULL,
        details_json TEXT NOT NULL
      );
      CREATE TABLE IF NOT EXISTS review_versions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        workspace_id TEXT NOT NULL,
        review_id TEXT NOT NULL,
        version INTEGER NOT NULL,
        content_hash TEXT NOT NULL,
        content TEXT NOT NULL,
        created_at TEXT NOT NULL,
        UNIQUE(workspace_id, review_id, version)
      );
      CREATE INDEX IF NOT EXISTS audit_by_workspace ON audit_log(workspace_id, id DESC);
      CREATE INDEX IF NOT EXISTS versions_by_review ON review_versions(workspace_id, review_id, version DESC);
      CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        email TEXT NOT NULL UNIQUE,
        display_name TEXT NOT NULL,
        password_hash TEXT NOT NULL,
        active INTEGER NOT NULL DEFAULT 1,
        failed_logins INTEGER NOT NULL DEFAULT 0,
        locked_until TEXT,
        created_at TEXT NOT NULL
      );
      CREATE TABLE IF NOT EXISTS workspaces (
        id TEXT PRIMARY KEY,
        name TEXT NOT NULL,
        created_at TEXT NOT NULL
      );
      CREATE TABLE IF NOT EXISTS memberships (
        workspace_id TEXT NOT NULL,
        user_id INTEGER NOT NULL,
        role TEXT NOT NULL,
        created_at TEXT NOT NULL,
        PRIMARY KEY(workspace_id,user_id),
        FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE,
        FOREIGN KEY(workspace_id) REFERENCES workspaces(id) ON DELETE CASCADE
      );
      CREATE TABLE IF NOT EXISTS sessions (
        token_hash TEXT PRIMARY KEY,
        user_id INTEGER NOT NULL,
        workspace_id TEXT NOT NULL,
        csrf_token TEXT NOT NULL,
        expires_at TEXT NOT NULL,
        created_at TEXT NOT NULL,
        FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE,
        FOREIGN KEY(workspace_id) REFERENCES workspaces(id) ON DELETE CASCADE
      );
      CREATE INDEX IF NOT EXISTS sessions_expiry ON sessions(expires_at);
      CREATE TABLE IF NOT EXISTS auth_events (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        email_hash TEXT NOT NULL,
        source_hash TEXT NOT NULL DEFAULT '',
        created_at TEXT NOT NULL,
        succeeded INTEGER NOT NULL
      );
      CREATE INDEX IF NOT EXISTS auth_events_recent ON auth_events(email_hash,created_at DESC);
    """)
    columns={row[1] for row in db.execute("PRAGMA table_info(auth_events)")}
    if "source_hash" not in columns:
        db.execute("ALTER TABLE auth_events ADD COLUMN source_hash TEXT NOT NULL DEFAULT ''")
    db.execute("CREATE INDEX IF NOT EXISTS auth_source_recent ON auth_events(source_hash,created_at DESC)")
    return db


@contextmanager
def db_session():
    db = connect()
    try:
        with db:
            yield db
    finally:
        db.close()


def valid_workspace(value: str | None) -> str:
    candidate = (value or DEFAULT_WORKSPACE).strip().lower()
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{0,63}", candidate):
        raise ValueError("Invalid workspace identifier")
    return candidate


class ServiceError(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status


def user_count() -> int:
    with db_session() as db:
        return int(db.execute("SELECT COUNT(*) FROM users").fetchone()[0])


def create_initial_admin(email: str, display_name: str, password: str, workspace_name: str = "Northstar Learning") -> None:
    email = email.strip().lower()
    if not re.fullmatch(r"[^\s@]{1,64}@[^\s@]{1,190}\.[^\s@]{2,24}", email):
        raise ValueError("Enter a valid email address")
    if not display_name.strip() or len(display_name.strip()) > 120:
        raise ValueError("Name must be between 1 and 120 characters")
    if not workspace_name.strip() or len(workspace_name.strip()) > 120:
        raise ValueError("Workspace name must be between 1 and 120 characters")
    password_hash = security.hash_password(password)
    workspace_id = "ws-" + secrets.token_hex(8)
    now = utc_now()
    with db_session() as db:
        db.execute("BEGIN IMMEDIATE")
        if db.execute("SELECT 1 FROM users LIMIT 1").fetchone():
            raise ValueError("An administrator is already configured")
        cur = db.execute("INSERT INTO users(email,display_name,password_hash,created_at) VALUES(?,?,?,?)", (email, display_name.strip(), password_hash, now))
        db.execute("INSERT INTO workspaces(id,name,created_at) VALUES(?,?,?)", (workspace_id, workspace_name.strip(), now))
        db.execute("INSERT INTO memberships(workspace_id,user_id,role,created_at) VALUES(?,?,?,?)", (workspace_id, cur.lastrowid, "Administrator", now))
        db.execute("INSERT INTO audit_log(workspace_id,created_at,actor,action,details_json) VALUES(?,?,?,?,?)", (workspace_id, now, display_name.strip(), "Initial administrator and workspace created", json.dumps({"text": "Initial administrator and workspace created", "actor": display_name.strip(), "role": "Administrator", "time": now}, ensure_ascii=False)))


def create_member(workspace: str, email: str, display_name: str, password: str, role: str) -> dict:
    allowed = {"Brand manager", "Content creator", "Legal reviewer", "Regional lead", "Administrator"}
    if role not in allowed:
        raise ValueError("Unsupported workspace role")
    email = email.strip().lower()
    if not re.fullmatch(r"[^\s@]{1,64}@[^\s@]{1,190}\.[^\s@]{2,24}", email):
        raise ValueError("Enter a valid email address")
    if not display_name.strip() or len(display_name.strip()) > 120:
        raise ValueError("Name must be between 1 and 120 characters")
    password_hash = security.hash_password(password)
    now = utc_now()
    with db_session() as db:
        if not db.execute("SELECT 1 FROM workspaces WHERE id=?", (workspace,)).fetchone():
            raise ServiceError(404, "Workspace not found")
        user = db.execute("SELECT id,display_name FROM users WHERE email=?", (email,)).fetchone()
        if user:
            user_id = user["id"]
            if db.execute("SELECT 1 FROM memberships WHERE user_id=? AND workspace_id=?", (user_id, workspace)).fetchone():
                raise ValueError("This user already belongs to the workspace")
            # Existing credentials belong to the user. A workspace administrator must not
            # be able to reset a global account password by adding it to another workspace.
        else:
            user_id = db.execute("INSERT INTO users(email,display_name,password_hash,created_at) VALUES(?,?,?,?)", (email, display_name.strip(), password_hash, now)).lastrowid
        db.execute("INSERT INTO memberships(workspace_id,user_id,role,created_at) VALUES(?,?,?,?) ON CONFLICT(workspace_id,user_id) DO UPDATE SET role=excluded.role", (workspace, user_id, role, now))
        return {"id": user_id, "email": email, "name": display_name.strip(), "role": role}


def membership_list(user_id: int) -> list[dict]:
    with db_session() as db:
        rows = db.execute("SELECT w.id,w.name,m.role FROM memberships m JOIN workspaces w ON w.id=m.workspace_id WHERE m.user_id=? ORDER BY w.name", (user_id,)).fetchall()
        return [{"id": row["id"], "name": row["name"], "role": row["role"]} for row in rows]


def create_session(user_id: int, workspace: str) -> tuple[str, str, str]:
    session_token, csrf_token = security.new_token(), security.new_token()
    expires = (datetime.now(timezone.utc) + timedelta(hours=SESSION_HOURS)).isoformat(timespec="seconds")
    with db_session() as db:
        if not db.execute("SELECT 1 FROM memberships WHERE workspace_id=? AND user_id=?", (workspace, user_id)).fetchone():
            raise ServiceError(403, "Workspace access denied")
        db.execute("DELETE FROM sessions WHERE expires_at<?", (utc_now(),))
        db.execute("INSERT INTO sessions(token_hash,user_id,workspace_id,csrf_token,expires_at,created_at) VALUES(?,?,?,?,?,?)", (security.token_hash(session_token), user_id, workspace, csrf_token, expires, utc_now()))
        excess=db.execute("SELECT token_hash FROM sessions WHERE user_id=? ORDER BY rowid DESC LIMIT -1 OFFSET 10",(user_id,)).fetchall()
        if excess: db.executemany("DELETE FROM sessions WHERE token_hash=?",[(item["token_hash"],) for item in excess])
    return session_token, csrf_token, expires


def login_user(email: str, password: str, source: str | None = None) -> tuple[dict, str, str, str]:
    normalized = email.strip().lower()
    email_hash = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
    source_hash = hashlib.sha256(source.encode("utf-8")).hexdigest() if source else ""
    now_dt = datetime.now(timezone.utc)
    cutoff = (now_dt - timedelta(minutes=15)).isoformat(timespec="seconds")
    failure: ServiceError | None = None
    row = member = None
    with db_session() as db:
        failures = db.execute("SELECT COUNT(*) FROM auth_events WHERE email_hash=? AND succeeded=0 AND created_at>=?", (email_hash, cutoff)).fetchone()[0]
        source_failures = db.execute("SELECT COUNT(*) FROM auth_events WHERE source_hash=? AND succeeded=0 AND created_at>=?", (source_hash, cutoff)).fetchone()[0] if source_hash else 0
        db.execute("DELETE FROM auth_events WHERE created_at<?", ((now_dt-timedelta(days=30)).isoformat(timespec="seconds"),))
        user = db.execute("SELECT * FROM users WHERE email=? AND active=1", (normalized,)).fetchone()
        if failures >= 10 or source_failures >= 25 or (user and user["locked_until"] and user["locked_until"] > utc_now()):
            failure=ServiceError(429,"Too many sign-in attempts. Try again in 15 minutes.")
        else:
            stored_hash = user["password_hash"] if user else _dummy_password_hash()
            valid = security.verify_password(password, stored_hash)
            if not user or not valid:
                db.execute("INSERT INTO auth_events(email_hash,source_hash,created_at,succeeded) VALUES(?,?,?,0)", (email_hash,source_hash,utc_now()))
                if user:
                    failed = int(user["failed_logins"]) + 1
                    locked_until = (now_dt + timedelta(minutes=15)).isoformat(timespec="seconds") if failed >= 5 else None
                    db.execute("UPDATE users SET failed_logins=?,locked_until=? WHERE id=?", (failed, locked_until, user["id"]))
                failure=ServiceError(401,"Email or password is incorrect")
            else:
                db.execute("INSERT INTO auth_events(email_hash,source_hash,created_at,succeeded) VALUES(?,?,?,1)", (email_hash,source_hash,utc_now()))
                db.execute("UPDATE users SET failed_logins=0,locked_until=NULL WHERE id=?", (user["id"],))
                member = db.execute("SELECT workspace_id FROM memberships WHERE user_id=? ORDER BY created_at LIMIT 1", (user["id"],)).fetchone()
                if not member:
                    failure=ServiceError(403,"No active workspace membership")
                else:
                    row = dict(user)
    if failure:
        raise failure
    token, csrf, expires = create_session(row["id"], member["workspace_id"])
    return {"id": row["id"], "email": row["email"], "name": row["display_name"], "workspace_id": member["workspace_id"]}, token, csrf, expires


def _dummy_password_hash() -> str:
    global DUMMY_PASSWORD_HASH
    try:
        return DUMMY_PASSWORD_HASH
    except NameError:
        DUMMY_PASSWORD_HASH = security.hash_password("dummy-password-used-for-timing")
        return DUMMY_PASSWORD_HASH


def resolve_session(token: str | None) -> dict | None:
    if not token or len(token) > 256:
        return None
    token_digest = security.token_hash(token)
    with db_session() as db:
        row = db.execute("SELECT s.*,u.email,u.display_name,u.active,m.role FROM sessions s JOIN users u ON u.id=s.user_id JOIN memberships m ON m.workspace_id=s.workspace_id AND m.user_id=s.user_id WHERE s.token_hash=?", (token_digest,)).fetchone()
        if not row or not row["active"] or row["expires_at"] <= utc_now():
            if row:
                db.execute("DELETE FROM sessions WHERE token_hash=?", (token_digest,))
            return None
        workspaces = db.execute("SELECT w.id,w.name,m.role FROM memberships m JOIN workspaces w ON w.id=m.workspace_id WHERE m.user_id=? ORDER BY w.name", (row["user_id"],)).fetchall()
        return {"user_id": row["user_id"], "email": row["email"], "name": row["display_name"], "workspace_id": row["workspace_id"], "role": row["role"], "csrf_token": row["csrf_token"], "workspaces": [{"id": r["id"], "name": r["name"], "role": r["role"]} for r in workspaces]}


def switch_session_workspace(token: str, workspace: str) -> str:
    new_csrf = security.new_token()
    digest = security.token_hash(token)
    with db_session() as db:
        session = db.execute("SELECT user_id FROM sessions WHERE token_hash=? AND expires_at>?", (digest, utc_now())).fetchone()
        if not session or not db.execute("SELECT 1 FROM memberships WHERE user_id=? AND workspace_id=?", (session["user_id"], workspace)).fetchone():
            raise ServiceError(403, "Workspace access denied")
        db.execute("UPDATE sessions SET workspace_id=?,csrf_token=? WHERE token_hash=?", (workspace, new_csrf, digest))
    return new_csrf


def create_workspace_for_user(user_id: int, name: str) -> dict:
    name=name.strip()
    if not name or len(name)>120: raise ValueError("Workspace name must be between 1 and 120 characters")
    slug=re.sub(r"[^a-z0-9]+","-",name.lower()).strip("-")[:36] or "workspace"
    workspace=f"{slug}-{secrets.token_hex(4)}"
    now=utc_now()
    with db_session() as db:
        db.execute("INSERT INTO workspaces(id,name,created_at) VALUES(?,?,?)",(workspace,name,now))
        db.execute("INSERT INTO memberships(workspace_id,user_id,role,created_at) VALUES(?,?,?,?)",(workspace,user_id,"Administrator",now))
        event=_event("BrandSentinel","System",f"Workspace created: {name}","◇")
        db.execute("INSERT INTO audit_log(workspace_id,created_at,actor,action,details_json) VALUES(?,?,?,?,?)",(workspace,now,"BrandSentinel",event["text"],json.dumps(event,ensure_ascii=False)))
    return {"id":workspace,"name":name,"role":"Administrator"}


def workspace_users(workspace: str) -> list[dict]:
    with db_session() as db:
        rows=db.execute("SELECT u.id,u.email,u.display_name,m.role,u.active,u.created_at FROM memberships m JOIN users u ON u.id=m.user_id WHERE m.workspace_id=? ORDER BY u.display_name",(workspace,)).fetchall()
        return [{"id":r["id"],"email":r["email"],"name":r["display_name"],"role":r["role"],"active":bool(r["active"]),"createdAt":r["created_at"]} for r in rows]


def update_member_role(workspace: str, user_id: int, role: str) -> dict:
    allowed={"Brand manager","Content creator","Legal reviewer","Regional lead","Administrator"}
    if role not in allowed: raise ValueError("Unsupported workspace role")
    with db_session() as db:
        member=db.execute("SELECT role FROM memberships WHERE workspace_id=? AND user_id=?",(workspace,user_id)).fetchone()
        if not member: raise ServiceError(404,"Workspace member not found")
        if member["role"]=="Administrator" and role!="Administrator":
            admins=db.execute("SELECT COUNT(*) FROM memberships m JOIN users u ON u.id=m.user_id WHERE m.workspace_id=? AND m.role='Administrator' AND u.active=1",(workspace,)).fetchone()[0]
            if admins<=1: raise ServiceError(409,"The last active workspace administrator cannot be demoted")
        db.execute("UPDATE memberships SET role=? WHERE workspace_id=? AND user_id=?",(role,workspace,user_id))
    return next(item for item in workspace_users(workspace) if item["id"]==user_id)


def append_audit(workspace: str, actor: str, role: str, text: str, icon: str="✓") -> None:
    item=_event(actor,role,text,icon)
    now=utc_now()
    with LOCK, db_session() as db:
        db.execute("INSERT INTO audit_log(workspace_id,created_at,actor,action,details_json) VALUES(?,?,?,?,?)",(workspace,now,actor,item["text"],json.dumps(item,ensure_ascii=False)))
        row=db.execute("SELECT state_json FROM workspace_state WHERE workspace_id=?",(workspace,)).fetchone()
        if row:
            state=json.loads(decrypt_text_for_workspace(workspace,"state",row["state_json"]))
            rows=db.execute("SELECT details_json FROM audit_log WHERE workspace_id=? ORDER BY id DESC LIMIT 30",(workspace,)).fetchall()
            state["audit"]=[json.loads(r["details_json"]) for r in rows]
            encrypted=security.encrypt_text(get_master_key(),workspace,"state",json.dumps(state,ensure_ascii=False,separators=(",",":")))
            db.execute("UPDATE workspace_state SET state_json=?,updated_at=? WHERE workspace_id=?",(encrypted,now,workspace))


def normalize_state(data: object) -> dict:
    if not isinstance(data, dict):
        raise ValueError("State must be an object")
    reviews = data.get("reviews", [])
    rules = data.get("rules", [])
    audit = data.get("audit", [])
    if not all(isinstance(x, list) for x in (reviews, rules, audit)):
        raise ValueError("Reviews, rules, and audit must be lists")
    if len(reviews) > 5000 or len(rules) > 5000 or len(audit) > 10000:
        raise ValueError("Workspace exceeds demo limits")
    id_pattern = re.compile(r"[A-Za-z0-9_-]{1,80}\Z")
    if not all(isinstance(r, dict) for r in reviews + rules):
        raise ValueError("Reviews and knowledge entries must be objects")
    review_ids=[r.get("id") for r in reviews]
    rule_ids=[r.get("id") for r in rules]
    if not all(isinstance(value,str) and id_pattern.fullmatch(value) for value in review_ids) or len(set(review_ids)) != len(reviews):
        raise ValueError("Review IDs must be unique and use letters, numbers, hyphens or underscores")
    if not all(isinstance(value,str) and id_pattern.fullmatch(value) for value in rule_ids) or len(set(rule_ids)) != len(rules):
        raise ValueError("Knowledge entry IDs must be unique and use letters, numbers, hyphens or underscores")
    for review in reviews:
        if not isinstance(review, dict) or not isinstance(review.get("id"), str) or not isinstance(review.get("title"), str):
            raise ValueError("Every review needs an ID and title")
        review.setdefault("text", "")
        review.setdefault("type", "Other")
        review.setdefault("campaign", "")
        review.setdefault("owner", "")
        review.setdefault("date", utc_now())
        review.setdefault("status", "In review")
        review.setdefault("comments", [])
        if not isinstance(review["comments"],list) or len(review["comments"])>1000:
            raise ValueError("Review comments must be a list of at most 1000 items")
        if not all(isinstance(review.get(key),str) for key in ("title","text","type","campaign","owner","date","status")):
            raise ValueError("Review fields must contain text")
        if review["status"] not in {"Draft","Needs changes","Changes requested","In review","Approved"}:
            raise ValueError("Unsupported review status")
        if not isinstance(review.get("findings", []), list):
            raise ValueError("Review findings must be a list")
        if len(review["title"]) > 300 or len(str(review.get("text", ""))) > MAX_BODY:
            raise ValueError("Review title or content is too long")
        finding_ids=[]
        for finding in review.get("findings", []):
            if not isinstance(finding, dict) or not finding.get("source"):
                raise ValueError("Each finding must include its evidence source")
            finding.setdefault("kind","Other")
            finding.setdefault("severity","Low")
            finding.setdefault("confidence",0)
            finding.setdefault("quote","")
            finding.setdefault("reason","")
            finding.setdefault("suggestion","")
            finding.setdefault("status","Open")
            finding.setdefault("method","Rule-based")
            if not all(isinstance(finding.get(key),str) for key in ("kind","severity","quote","reason","source","suggestion","status","method")):
                raise ValueError("Finding fields must contain text")
            if finding["severity"] not in {"High","Medium","Low"}:
                raise ValueError("Finding severity must be High, Medium or Low")
            if finding["status"] not in {"Open","Accepted","Dismissed","Escalated","Reviewer cleared"}:
                raise ValueError("Unsupported finding status")
            finding_id=finding.get("id")
            if not isinstance(finding_id,str) or not id_pattern.fullmatch(finding_id):
                raise ValueError("Finding IDs must use letters, numbers, hyphens or underscores")
            finding_ids.append(finding_id)
        if len(set(finding_ids)) != len(finding_ids):
            raise ValueError("Finding IDs must be unique within a review")
        for comment in review["comments"]:
            if not isinstance(comment,dict) or not isinstance(comment.get("text"),str) or len(comment["text"])>10000:
                raise ValueError("Comments must contain text of at most 10000 characters")
        if review.get("status") == "Approved" and any(
            f.get("status") in {"Open", "Escalated"}
            or (f.get("kind") in {"Unsupported claim", "Unverified fact"} and f.get("status") != "Reviewer cleared")
            for f in review.get("findings", [])
        ):
            raise ValueError("Open, escalated, or uncleared high-risk findings block approval")
    for rule in rules:
        if not isinstance(rule.get("title"), str) or len(rule["title"]) > 300 or len(str(rule.get("detail", ""))) > 100_000:
            raise ValueError("Every knowledge entry needs a title")
        for key in ("detail","category","owner","status","source","bannedTerm","preferredTerm","reviewDate"):
            if key in rule and not isinstance(rule[key],str): raise ValueError("Knowledge entry fields must contain text")
        rule.setdefault("detail",""); rule.setdefault("category","Other"); rule.setdefault("status","Active")
    return {"rules": rules, "reviews": reviews, "audit": audit}


def _event(actor: str, role: str, text: str, icon: str = "✓") -> dict:
    return {"icon": icon, "text": text[:240], "time": "Just now", "actor": actor[:120], "role": role[:60]}


def server_findings(text: str, knowledge: list[dict]) -> list[dict]:
    """Re-run core deterministic checks on the trusted server before accepting state."""
    findings=[]
    def add(kind,severity,confidence,quote,reason,source,suggestion):
        findings.append({"id":"sf"+secrets.token_hex(6),"kind":kind,"severity":severity,"confidence":confidence,"quote":quote,"reason":reason,"source":source,"suggestion":suggestion,"status":"Open","method":"Rule-based"})
    absolute=re.search(r"\b(guarantee(?:s|d)?|100%|risk-free|no risk|always|never)\b[^.!?]*",text,re.I)
    if absolute:
        add("Unsupported claim","High",97,absolute[0].strip(),"Absolute outcome language needs an approved claim and supporting evidence. No matching approved claim is available in this knowledge base.","Claims library · No approved claim found","Replace with a precise, evidence-backed statement; ask the authorised reviewer if the claim is essential.")
    configured=[]
    for rule in knowledge:
        banned=rule.get("bannedTerm",""); preferred=rule.get("preferredTerm","")
        if rule.get("category")=="Terminology" and banned and preferred:
            for match in re.finditer(r"\b"+re.escape(banned)+r"\b",text,re.I):
                configured.append(match.start())
                add("Terminology","Medium",96,match[0],f"The approved terminology entry prefers “{preferred}”.",f"{rule.get('source') or 'Brand knowledge'} · {rule.get('title','')}",preferred)
    if not any(r.get("category")=="Terminology" and str(r.get("bannedTerm","")).lower()=="students" for r in knowledge):
        for match in re.finditer(r"\bstudents\b",text,re.I):
            if match.start() not in configured:
                add("Terminology","Medium",96,match[0],"The approved glossary prefers “learners” for people enrolled in a programme.","Brand guideline · Terminology rule","learners")
    link=re.search(r"\b(click here|learn more here)\b",text,re.I)
    if link:
        add("Accessibility","Low",99,link[0],"This link text does not describe its destination for screen-reader users.","Accessibility rule pack · Descriptive link text","Use descriptive link text, such as “Explore the programme”.")
    hype=re.search(r"\b(achieve your dreams|best in the world|revolutionary|once-in-a-lifetime)\b",text,re.I)
    if hype:
        add("Tone of voice","Low",88,hype[0],"This phrase may sound promotional rather than clear, practical and encouraging.","Brand guideline · Voice principles","Describe the practical benefit in plain language.")
    facts=" ".join(str(rule.get("detail","")).lower() for rule in knowledge if rule.get("category") in {"Approved facts","Claims"})
    for match in re.finditer(r"\b\d+(?:\.\d+)?\s?(?:%|percent|times?|x|weeks?|months?|years?)\b",text,re.I):
        if match[0].lower() in facts or any(item["kind"]=="Unsupported claim" and match[0] in item["quote"] for item in findings):
            continue
        start=text.rfind(".",0,match.start())+1; end=text.find(".",match.start())
        quote=text[start:len(text) if end<0 else end].strip() or match[0]
        add("Unverified fact","Medium",92,quote,"This measurable statement does not match an approved fact or claim in the knowledge base. A person should confirm the source before publication.","Approved facts register · No matching entry","Add the approved source or revise the statement to a verifiable fact.")
    return findings


def enforce_server_findings(old: dict, state: dict) -> None:
    previous={review["id"]:review for review in old.get("reviews",[])}
    for review in state["reviews"]:
        prior=previous.get(review["id"])
        if prior and prior.get("text")==review.get("text"):
            continue
        known={(finding.get("kind"),finding.get("quote")) for finding in review.get("findings",[]) if finding.get("status")=="Open"}
        for finding in server_findings(review.get("text", ""),state["rules"]):
            key=(finding["kind"],finding["quote"])
            if key in known:
                existing=next(item for item in review["findings"] if item.get("status")=="Open" and (item.get("kind"),item.get("quote"))==key)
                for field in ("severity","confidence","reason","source","suggestion","method"):
                    existing[field]=finding[field]
            else:
                review.setdefault("findings",[]).append(finding)
                known.add(key)


def authorize_transition(old: dict, new: dict, role: str, actor: str) -> list[dict]:
    privileged = {"Brand manager", "Administrator"}
    if not {r.get("id") for r in old["reviews"]}.issubset({r.get("id") for r in new["reviews"]}):
        raise ServiceError(403, "Reviews cannot be removed from the audit history")
    if not {r.get("id") for r in old["rules"]}.issubset({r.get("id") for r in new["rules"]}):
        raise ServiceError(403, "Knowledge entries cannot be deleted from this workflow")
    old_reviews = {r["id"]: r for r in old["reviews"]}
    new_reviews = {r["id"]: r for r in new["reviews"]}
    events=[]
    for review in new["reviews"]:
        previous=old_reviews.get(review["id"])
        if previous is None:
            if role not in {"Content creator", *privileged}:
                raise ServiceError(403, "This role cannot submit content")
            if review.get("status")=="Approved":
                raise ServiceError(403,"New content cannot be created in an approved state")
            if any(f.get("status")!="Open" for f in review.get("findings",[])):
                raise ServiceError(403,"New findings must start open")
            if review.get("comments"):
                raise ServiceError(403,"New reviews cannot include pre-created comments")
            review["owner"] = actor
            review["resubmissionCount"] = 0
            review.pop("approvedAt",None)
            events.append(_event(actor, role, f"Submitted review {review['id']}: {review['title']}", "↗"))
            continue
        if previous.get("status")=="Approved":
            immutable_before={key:value for key,value in previous.items() if key!="comments"}
            immutable_after={key:value for key,value in review.items() if key!="comments"}
            if immutable_before!=immutable_after:
                raise ServiceError(409, "Approved content and decisions are locked")
        if review.get("text")!=previous.get("text"):
            if role not in {"Content creator", *privileged}:
                raise ServiceError(403, "This role cannot edit submitted content")
            events.append(_event(actor, role, f"Resubmitted review {review['id']}: {review['title']}", "↗"))
            review["resubmissionCount"]=int(previous.get("resubmissionCount",0))+1
        else:
            if review.get("resubmissionCount",0)!=previous.get("resubmissionCount",0):
                raise ServiceError(403,"Review version metadata is server managed")
        if review.get("owner")!=previous.get("owner"):
            raise ServiceError(403,"Review ownership is immutable")
        allowed_status={
            "Draft":{"Draft","In review","Needs changes"},
            "Needs changes":{"Needs changes","In review","Changes requested","Approved"},
            "Changes requested":{"Changes requested","Needs changes","In review","Approved"},
            "In review":{"In review","Needs changes","Changes requested","Approved"},
        }
        if review.get("status")!=previous.get("status") and review.get("status") not in allowed_status.get(previous.get("status"),set()):
            raise ServiceError(403,"Invalid review status transition")
        old_findings={f.get("id"):f for f in previous.get("findings",[])}
        current_findings={f.get("id"):f for f in review.get("findings",[])}
        if len(current_findings)<len(old_findings) or not set(old_findings).issubset(current_findings):
            raise ServiceError(403, "Findings cannot be removed from a review")
        for fid, old_f in old_findings.items():
            finding=current_findings[fid]
            for key in ("kind","quote","source","reason"):
                if finding.get(key)!=old_f.get(key):
                    raise ServiceError(403, "Finding evidence is immutable")
            before,after=old_f.get("status"),finding.get("status")
            if before==after:
                for key in ("note","creatorDecision","creatorNote","assignedTo"):
                    if finding.get(key)!=old_f.get(key):
                        raise ServiceError(403,"Finding decision metadata cannot be edited after a decision")
                continue
            highrisk=old_f.get("kind") in {"Unsupported claim","Unverified fact"}
            if after in {"Accepted","Dismissed","Escalated"}:
                if role not in {"Content creator", *privileged}:
                    raise ServiceError(403, "This role cannot decide this finding")
                if after in {"Dismissed","Escalated"} and not str(finding.get("note","")).strip():
                    raise ValueError("Dismissal and escalation require a written reason")
            elif after=="Reviewer cleared":
                allowed=(role in {"Legal reviewer","Administrator"} if highrisk else
                         role in {"Regional lead","Brand manager","Administrator"} if old_f.get("kind")=="Terminology" else role in privileged)
                if not allowed or not str(finding.get("note","")).strip():
                    raise ServiceError(403, "An authorised reviewer and written rationale are required")
            elif after=="Open" and before=="Accepted" and role in {"Content creator", *privileged}:
                pass
            else:
                raise ServiceError(403, "Invalid finding state transition")
            events.append(_event(actor,role,f"{after} {old_f.get('kind','finding')} in review {review['id']}: {review['title']}"))
        old_comments=previous.get("comments",[])
        comments=review.get("comments",[])
        if len(comments)<len(old_comments) or comments[:len(old_comments)]!=old_comments:
            raise ServiceError(403,"Review comments cannot be edited or removed")
        if len(comments)>len(old_comments):
            if len(comments)>len(old_comments)+20:
                raise ValueError("Too many comments in one update")
            for comment in comments[len(old_comments):]:
                if not isinstance(comment,dict) or not str(comment.get("text","")).strip() or len(comment.get("text",""))>10000:
                    raise ValueError("Comments must contain 1 to 10000 characters")
                comment["author"]=actor
                comment["createdAt"]=utc_now()
                events.append(_event(actor,role,f"Commented on review {review['id']}: {review['title']}","◌"))
        if review.get("status")=="Approved" and previous.get("status")!="Approved":
            if role not in privileged:
                raise ServiceError(403,"Only a brand manager or administrator can approve")
            if any(f.get("status") in {"Open","Escalated"} or (f.get("kind") in {"Unsupported claim","Unverified fact"} and f.get("status")!="Reviewer cleared") for f in review.get("findings",[])):
                raise ValueError("Open, escalated, or uncleared high-risk findings block approval")
            review["approvedAt"]=utc_now()
            events.append(_event(actor,role,f"Approved review {review['id']}: {review['title']}"))
        elif previous.get("approvedAt")!=review.get("approvedAt"):
            raise ServiceError(403,"Approval timestamps are server managed")
    old_rules={r.get("id"):r for r in old["rules"]}
    new_rules={r.get("id"):r for r in new["rules"]}
    for rid, rule in new_rules.items():
        if old_rules.get(rid)!=rule:
            if role not in privileged:
                raise ServiceError(403,"Only a brand manager or administrator can change approved knowledge")
            events.append(_event(actor,role,f"Updated knowledge entry: {str(rule.get('title',''))[:120]}","◇"))
    return events


def save_state(workspace: str, payload: dict, actor: str = "Local workspace user", role: str = "Brand manager") -> None:
    expected_revision=payload.get("_revision") if isinstance(payload,dict) else None
    if expected_revision is not None and (not isinstance(expected_revision,str) or (expected_revision and not re.fullmatch(r"[a-f0-9]{64}",expected_revision))):
        raise ValueError("Invalid workspace revision")
    state = normalize_state(payload)
    import hashlib
    now = utc_now()
    with LOCK, db_session() as db:
        old_row = db.execute("SELECT state_json FROM workspace_state WHERE workspace_id=?", (workspace,)).fetchone()
        actual_revision=hashlib.sha256(old_row["state_json"].encode("utf-8")).hexdigest() if old_row else ""
        if expected_revision is not None and expected_revision!=actual_revision:
            raise ServiceError(409,"This workspace changed in another session. Reload the latest version and reapply your change.")
        old_state = json.loads(decrypt_text_for_workspace(workspace,"state",old_row["state_json"])) if old_row else {"reviews": [], "rules": [], "audit": []}
        enforce_server_findings(old_state,state)
        if len(json.dumps(state,ensure_ascii=False,separators=(",",":" )).encode("utf-8"))>MAX_BODY:
            raise ValueError("Workspace state is too large")
        for review in state["reviews"]:
            if review["id"] not in {item["id"] for item in old_state.get("reviews",[])} and review["comments"]:
                raise ServiceError(403,"New reviews cannot include pre-created comments")
        generated_events=authorize_transition(old_state,state,role,actor)
        old_reviews = {r["id"]: r for r in old_state.get("reviews", [])}
        for review in state["reviews"]:
            content = str(review.get("text", ""))
            previous = old_reviews.get(review["id"])
            if previous and previous.get("status") == "Approved" and previous.get("text") != content:
                raise ValueError("Approved content is locked and cannot be edited")
            if previous is None or previous.get("text") != content:
                row = db.execute("SELECT COALESCE(MAX(version),0) AS n FROM review_versions WHERE workspace_id=? AND review_id=?", (workspace, review["id"])).fetchone()
                if row["n"]>=100:
                    raise ValueError("A review reached the 100-version retention limit")
                version = row["n"] + 1
                digest = hashlib.sha256(content.encode("utf-8")).hexdigest()
                encrypted_content=security.encrypt_text(get_master_key(),workspace,f"review:{review['id']}:v{version}",content)
                db.execute("INSERT INTO review_versions(workspace_id,review_id,version,content_hash,content,created_at) VALUES(?,?,?,?,?,?)", (workspace, review["id"], version, digest, encrypted_content, now))
        for item in generated_events:
            db.execute("INSERT INTO audit_log(workspace_id,created_at,actor,action,details_json) VALUES(?,?,?,?,?)", (workspace, now, item["actor"], item["text"], json.dumps(item, ensure_ascii=False)))
        audit_rows=db.execute("SELECT details_json FROM audit_log WHERE workspace_id=? ORDER BY id DESC LIMIT 30",(workspace,)).fetchall()
        state["audit"]=[json.loads(row["details_json"]) for row in audit_rows]
        state_json = json.dumps(state, ensure_ascii=False, separators=(",", ":"))
        encrypted_state=security.encrypt_text(get_master_key(),workspace,"state",json.dumps(state,ensure_ascii=False,separators=(",",":")))
        db.execute("INSERT INTO workspace_state(workspace_id,state_json,updated_at) VALUES(?,?,?) ON CONFLICT(workspace_id) DO UPDATE SET state_json=excluded.state_json,updated_at=excluded.updated_at", (workspace, encrypted_state, now))


def load_state(workspace: str) -> dict:
    with LOCK, db_session() as db:
        row = db.execute("SELECT state_json,updated_at FROM workspace_state WHERE workspace_id=?", (workspace,)).fetchone()
        if not row: return {"rules": [], "reviews": [], "audit": [], "_revision":""}
        state=json.loads(decrypt_text_for_workspace(workspace,"state",row["state_json"]))
        if not row["state_json"].startswith("v1."):
            migrated=security.encrypt_text(get_master_key(),workspace,"state",json.dumps(state,ensure_ascii=False,separators=(",",":")))
            db.execute("UPDATE workspace_state SET state_json=? WHERE workspace_id=?",(migrated,workspace))
            row_state=migrated
        else: row_state=row["state_json"]
        state["_revision"]=hashlib.sha256(row_state.encode("utf-8")).hexdigest()
        rows=db.execute("SELECT details_json FROM audit_log WHERE workspace_id=? ORDER BY id DESC LIMIT 30",(workspace,)).fetchall()
        state["audit"]=[json.loads(item["details_json"]) for item in rows]
        return state


def audit_entries(workspace: str, limit: int = 200) -> list[dict]:
    with LOCK, db_session() as db:
        rows = db.execute("SELECT created_at,actor,action,details_json FROM audit_log WHERE workspace_id=? ORDER BY id DESC LIMIT ?", (workspace, min(max(limit, 1), 500))).fetchall()
        return [{"createdAt": row["created_at"], "actor": row["actor"], "action": row["action"], "details": json.loads(row["details_json"])} for row in rows]


def review_versions(workspace: str, review_id: str) -> list[dict]:
    with LOCK, db_session() as db:
        rows = db.execute("SELECT version,content_hash,content,created_at FROM review_versions WHERE workspace_id=? AND review_id=? ORDER BY version DESC", (workspace, review_id)).fetchall()
        output=[]
        for row in rows:
            purpose=f"review:{review_id}:v{row['version']}"
            content=decrypt_text_for_workspace(workspace,purpose,row["content"])
            if not row["content"].startswith("v1."):
                encrypted=security.encrypt_text(get_master_key(),workspace,purpose,content)
                db.execute("UPDATE review_versions SET content=? WHERE workspace_id=? AND review_id=? AND version=?",(encrypted,workspace,review_id,row["version"]))
            output.append({"version":row["version"],"content_hash":row["content_hash"],"content":content,"created_at":row["created_at"]})
        return output


def decrypt_text_for_workspace(workspace: str, purpose: str, value: str) -> str:
    if value.startswith("v1."):
        return security.decrypt_text(get_master_key(),workspace,purpose,value)
    # One-time migration path for pre-encryption local demo databases.
    return value


def extract_text(filename: str, raw: bytes) -> dict:
    """Extract text from supported office/web formats; never executes embedded content."""
    from html.parser import HTMLParser
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        raise ValueError("Supported formats: TXT, Markdown, HTML, PDF, DOCX and PPTX")
    if len(raw) > 15 * 1024 * 1024:
        raise ValueError("File exceeds the 15 MB limit")
    if ext in {".docx", ".pptx"}:
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries=archive.infolist()
                total=sum(item.file_size for item in entries)
                if len(entries)>2000 or total>100*1024*1024:
                    raise ValueError("Office document exceeds safe parsing limits")
                if any(item.filename.startswith(("/", "\\")) or ".." in Path(item.filename).parts for item in entries):
                    raise ValueError("Office document contains an unsafe path")
                if any(item.compress_size and item.file_size/item.compress_size>1000 for item in entries):
                    raise ValueError("Office document has an unsafe compression ratio")
        except zipfile.BadZipFile as exc:
            raise ValueError("Office file is not a valid DOCX or PPTX archive") from exc
    checks=[]
    if ext in {".txt", ".md"}:
        text = raw.decode("utf-8-sig", errors="replace")
    elif ext in {".html", ".htm"}:
        class TextOnly(HTMLParser):
            def __init__(self):
                super().__init__(); self.parts=[]; self.hidden=0; self.links=[]; self.previous_heading=0
            def handle_starttag(self, tag, attrs):
                attrs=dict(attrs)
                if tag in {"script", "style", "noscript"}: self.hidden += 1
                elif tag in {"p", "br", "div", "li", "h1", "h2", "h3", "tr"}: self.parts.append("\n")
                if tag=="html" and not attrs.get("lang"):
                    checks.append({"kind":"Accessibility","severity":"Medium","confidence":100,"quote":"<html>","reason":"The document does not declare a primary language, which helps assistive technology pronounce its text.","source":"Accessibility rule pack · Document language","suggestion":"Set a valid lang attribute on the html element, such as lang=\"en\"."})
                if tag=="img" and not attrs.get("alt"):
                    checks.append({"kind":"Accessibility","severity":"High","confidence":100,"quote":"<img>","reason":"An image has no alternative text. Users of screen readers may miss its meaning.","source":"Accessibility rule pack · Image alternative text","suggestion":"Add concise alt text that communicates the image's purpose."})
                if tag=="a": self.links.append({"attrs":attrs,"text":[]})
                if re.fullmatch(r"h[1-6]",tag):
                    level=int(tag[1])
                    if self.previous_heading and level>self.previous_heading+1:
                        checks.append({"kind":"Accessibility","severity":"Medium","confidence":100,"quote":f"<{tag}>","reason":"Heading levels skip a level, which makes page structure harder to navigate.","source":"Accessibility rule pack · Heading order","suggestion":f"Use an h{self.previous_heading+1} heading before this section, or adjust the heading level."})
                    self.previous_heading=level
            def handle_endtag(self, tag):
                if tag in {"script", "style", "noscript"} and self.hidden: self.hidden -= 1
                elif tag in {"p", "div", "li", "h1", "h2", "h3", "tr"}: self.parts.append("\n")
                if tag=="a" and self.links:
                    link=self.links.pop(); label=" ".join(link["text"]).strip(); aria=link["attrs"].get("aria-label","")
                    if not label and not aria or re.fullmatch(r"(?:click here|here|learn more here)",label,re.I):
                        checks.append({"kind":"Accessibility","severity":"Low","confidence":100,"quote":label or "<a>","reason":"This link does not describe its destination for screen-reader users.","source":"Accessibility rule pack · Descriptive link text","suggestion":"Use link text that describes the destination or action."})
            def handle_data(self, data):
                if not self.hidden:
                    self.parts.append(data)
                    for link in self.links: link["text"].append(data)
        parser=TextOnly(); parser.feed(raw.decode("utf-8",errors="replace")); text="".join(parser.parts)
    elif ext == ".pdf":
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(raw), strict=False)
        if len(reader.pages) > 100: raise ValueError("PDF has more than 100 pages")
        text="\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        from docx import Document
        doc=Document(io.BytesIO(raw))
        text="\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
    else:
        from pptx import Presentation
        deck=Presentation(io.BytesIO(raw))
        if len(deck.slides) > 200: raise ValueError("Presentation has more than 200 slides")
        text="\n".join(shape.text for slide in deck.slides for shape in slide.shapes if getattr(shape,"has_text_frame",False))
    text=re.sub(r"\n{3,}","\n\n",text).strip()
    if not text: raise ValueError("No selectable text found. Scanned files and OCR are not supported in this demo.")
    if len(text) > 2_000_000: text=text[:2_000_000]
    return {"filename":Path(filename).name[:255],"format":ext[1:].upper(),"text":text,"characters":len(text),"truncated":len(text)>=2_000_000,"checks":checks}


def parse_multipart(content_type: str, body: bytes) -> tuple[str, bytes]:
    headers = f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("ascii", errors="replace") + body
    message=BytesParser(policy=policy.default).parsebytes(headers)
    if not message.is_multipart(): raise ValueError("Expected a multipart upload")
    for part in message.iter_parts():
        filename=part.get_filename()
        if filename: return filename, part.get_payload(decode=True) or b""
    raise ValueError("No file was included")


class SecureHandler(BaseHTTPRequestHandler):
    server_version = "BrandSentinel"
    sys_version = ""

    def setup(self):
        super().setup()
        self.request.settimeout(30)

    def log_message(self, fmt, *args):
        # Only the path and status are logged. Request bodies, headers and query strings are excluded.
        path=self.path.split("?",1)[0]
        print(f"{self.log_date_time_string()} {self.command} {path[:160]} {args[1] if len(args)>1 else ''}")

    def send_json(self, status: int, value: object, extra_headers: list[tuple[str,str]] | None=None):
        raw=json.dumps(value,ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type","application/json; charset=utf-8")
        self.send_header("Content-Length",str(len(raw)))
        self.send_header("Cache-Control","no-store")
        self.send_header("X-Content-Type-Options","nosniff")
        self.send_header("Referrer-Policy","no-referrer")
        self.send_header("X-Frame-Options","DENY")
        self.send_header("Permissions-Policy","camera=(), microphone=(), geolocation=()")
        if os.environ.get("BRANDSENTINEL_ENV")=="production": self.send_header("Strict-Transport-Security","max-age=31536000; includeSubDomains")
        for name,value2 in extra_headers or []: self.send_header(name,value2)
        self.end_headers()
        if self.command!="HEAD": self.wfile.write(raw)

    def send_error_json(self, status: int, message: str):
        self.send_json(status,{"error":message})

    def read_body(self) -> bytes:
        try: length=int(self.headers.get("Content-Length","-1"))
        except ValueError: raise ValueError("Invalid Content-Length")
        if length<0 or length>MAX_BODY: raise ServiceError(413,"Request body exceeds the 16 MB limit")
        body=self.rfile.read(length)
        if len(body)!=length: raise ValueError("Incomplete request body")
        return body

    def read_json(self) -> dict:
        if self.headers.get_content_type()!="application/json": raise ServiceError(415,"Content-Type must be application/json")
        try: payload=json.loads(self.read_body().decode("utf-8"))
        except (UnicodeDecodeError,json.JSONDecodeError) as exc: raise ValueError("Request body must contain valid UTF-8 JSON") from exc
        if not isinstance(payload,dict): raise ValueError("Request body must be a JSON object")
        return payload

    def cookie_token(self) -> str | None:
        cookie=SimpleCookie()
        try: cookie.load(self.headers.get("Cookie",""))
        except Exception: return None
        morsel=cookie.get(cookie_name())
        return morsel.value if morsel else None

    def require_session(self) -> dict:
        current=resolve_session(self.cookie_token())
        if not current: raise ServiceError(401,"Sign in to continue")
        return current

    def require_csrf(self, current: dict) -> None:
        self.verify_origin()
        supplied=self.headers.get("X-CSRF-Token","")
        if not supplied or not hmac.compare_digest(supplied,current["csrf_token"]):
            raise ServiceError(403,"CSRF validation failed; reload the page and try again")

    def validate_host(self) -> None:
        host=self.headers.get("Host","").lower()
        allowed={item.strip().lower() for item in os.environ.get("BRANDSENTINEL_ALLOWED_HOSTS","").split(",") if item.strip()}
        if os.environ.get("BRANDSENTINEL_ENV")=="production" and (not allowed or host not in allowed):
            raise ServiceError(400,"Host is not configured for this deployment")

    def verify_origin(self) -> None:
        site=self.headers.get("Sec-Fetch-Site","").lower()
        if site=="cross-site": raise ServiceError(403,"Cross-site requests are not allowed")
        self.validate_host()
        origin=self.headers.get("Origin")
        if origin:
            parsed=urlparse(origin)
            if not parsed.netloc or parsed.netloc.lower()!=host:
                raise ServiceError(403,"Request origin is not allowed")
            if os.environ.get("BRANDSENTINEL_ENV")=="production" and parsed.scheme!="https":
                raise ServiceError(403,"HTTPS origin is required")

    def session_cookie(self, token: str, max_age: int) -> str:
        secure=os.environ.get("BRANDSENTINEL_ENV")=="production" or os.environ.get("BRANDSENTINEL_COOKIE_SECURE")=="1"
        parts=[f"{cookie_name()}={token}","Path=/","HttpOnly","SameSite=Strict",f"Max-Age={max_age}"]
        if secure: parts.append("Secure")
        return "; ".join(parts)

    def active_workspace(self, current: dict) -> dict:
        return next((item for item in current["workspaces"] if item["id"]==current["workspace_id"]),{})

    def static_response(self, path: str):
        rel="index.html" if path=="/" else path.lstrip("/")
        if rel not in STATIC_FILES: return self.send_error_json(404,"Not found")
        target=FRONTEND_DIR/rel
        mime={".html":"text/html; charset=utf-8",".css":"text/css; charset=utf-8",".js":"text/javascript; charset=utf-8"}.get(target.suffix,"application/octet-stream")
        raw=target.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type",mime)
        self.send_header("Content-Length",str(len(raw)))
        self.send_header("Cache-Control","no-cache")
        self.send_header("X-Content-Type-Options","nosniff")
        self.send_header("X-Frame-Options","DENY")
        self.send_header("Referrer-Policy","strict-origin-when-cross-origin")
        self.send_header("Content-Security-Policy","default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; object-src 'none'; base-uri 'self'; frame-ancestors 'none'; form-action 'self'")
        if os.environ.get("BRANDSENTINEL_ENV")=="production": self.send_header("Strict-Transport-Security","max-age=31536000; includeSubDomains")
        self.end_headers()
        if self.command!="HEAD": self.wfile.write(raw)

    def do_GET(self):
        try:
            path=urlparse(self.path).path
            if path!="/healthz": self.validate_host()
            if path=="/api/auth/status":
                return self.send_json(200,{"setupRequired":user_count()==0})
            if path=="/healthz":
                return self.send_json(200,{"status":"ok"})
            if path=="/api/auth/me":
                current=self.require_session(); workspace=self.active_workspace(current)
                return self.send_json(200,{"user":{"id":current["user_id"],"email":current["email"],"name":current["name"]},"role":current["role"],"workspace":{"id":current["workspace_id"],"name":workspace.get("name","Workspace")},"workspaces":current["workspaces"],"csrfToken":current["csrf_token"]})
            if path=="/api/state":
                current=self.require_session(); return self.send_json(200,load_state(current["workspace_id"]))
            if path=="/api/audit":
                current=self.require_session(); return self.send_json(200,{"entries":audit_entries(current["workspace_id"],500)})
            if path=="/api/users":
                current=self.require_session()
                if current["role"]!="Administrator": raise ServiceError(403,"Administrator role required")
                return self.send_json(200,{"users":workspace_users(current["workspace_id"])})
            match=re.fullmatch(r"/api/reviews/([A-Za-z0-9_-]+)/versions",path)
            if match:
                current=self.require_session(); return self.send_json(200,{"versions":review_versions(current["workspace_id"],match.group(1))})
            if path.startswith("/api/"): return self.send_error_json(404,"API route not found")
            return self.static_response(path)
        except ServiceError as exc: self.send_error_json(exc.status,str(exc))
        except ValueError as exc: self.send_error_json(400,str(exc))
        except Exception as exc: self.send_error_json(500,"Internal server error"); print(f"GET failure: {type(exc).__name__}")

    def do_POST(self):
        try:
            path=urlparse(self.path).path
            self.verify_origin()
            if path=="/api/auth/login":
                payload=self.read_json(); email=str(payload.get("email",""))[:256]; password=str(payload.get("password",""))
                if len(password)>1024: raise ValueError("Invalid credentials")
                user,token,csrf,expires=login_user(email,password,request_source(self))
                current=resolve_session(token); workspace=self.active_workspace(current)
                cookie=self.session_cookie(token,SESSION_HOURS*3600)
                return self.send_json(200,{"user":{"id":user["id"],"email":user["email"],"name":user["name"]},"role":current["role"],"workspace":{"id":current["workspace_id"],"name":workspace.get("name","Workspace")},"workspaces":current["workspaces"],"csrfToken":csrf},[("Set-Cookie",cookie)])
            current=self.require_session(); self.require_csrf(current)
            if path=="/api/auth/logout":
                token=self.cookie_token()
                with db_session() as db: db.execute("DELETE FROM sessions WHERE token_hash=?",(security.token_hash(token),))
                return self.send_json(200,{"ok":True},[("Set-Cookie",self.session_cookie("",0))])
            if path=="/api/workspaces/switch":
                payload=self.read_json(); workspace=valid_workspace(str(payload.get("workspaceId","")))
                csrf=switch_session_workspace(self.cookie_token(),workspace)
                updated=resolve_session(self.cookie_token()); selected=self.active_workspace(updated)
                append_audit(workspace,updated["name"],updated["role"],f"Signed in to workspace {selected.get('name','Workspace')}","↔")
                return self.send_json(200,{"workspace":{"id":workspace,"name":selected.get("name")},"role":updated["role"],"workspaces":updated["workspaces"],"csrfToken":csrf})
            if path=="/api/users":
                if current["role"]!="Administrator": raise ServiceError(403,"Administrator role required")
                payload=self.read_json(); member=create_member(current["workspace_id"],str(payload.get("email","")),str(payload.get("name","")),str(payload.get("password","")),str(payload.get("role","")))
                append_audit(current["workspace_id"],current["name"],current["role"],f"Added workspace member {member['email']} as {member['role']}","＋")
                return self.send_json(201,{"user":member})
            if path=="/api/workspaces":
                if current["role"]!="Administrator": raise ServiceError(403,"Administrator role required")
                payload=self.read_json(); workspace=create_workspace_for_user(current["user_id"],str(payload.get("name","")))
                return self.send_json(201,{"workspace":workspace,"workspaces":membership_list(current["user_id"])})
            if path=="/api/extract":
                content_type=self.headers.get("Content-Type","")
                filename,data=parse_multipart(content_type,self.read_body())
                return self.send_json(200,extract_text(filename,data))
            return self.send_error_json(404,"API route not found")
        except ServiceError as exc: self.send_error_json(exc.status,str(exc))
        except ValueError as exc: self.send_error_json(400,str(exc))
        except Exception as exc: self.send_error_json(422,"Request could not be processed"); print(f"POST failure: {type(exc).__name__}")

    def do_PUT(self):
        try:
            path=urlparse(self.path).path
            self.verify_origin(); current=self.require_session(); self.require_csrf(current)
            member_match=re.fullmatch(r"/api/users/(\d+)",path)
            if member_match:
                if current["role"]!="Administrator": raise ServiceError(403,"Administrator role required")
                payload=self.read_json(); member=update_member_role(current["workspace_id"],int(member_match.group(1)),str(payload.get("role","")))
                append_audit(current["workspace_id"],current["name"],current["role"],f"Changed {member['email']} role to {member['role']}","♙")
                return self.send_json(200,{"user":member})
            if path!="/api/state": return self.send_error_json(404,"API route not found")
            payload=self.read_json()
            if not isinstance(payload.get("_revision"),str): raise ServiceError(428,"Reload the workspace before saving changes")
            save_state(current["workspace_id"],payload,current["name"],current["role"])
            return self.send_json(200,{"ok":True,"state":load_state(current["workspace_id"]),"updatedAt":utc_now()})
        except ServiceError as exc: self.send_error_json(exc.status,str(exc))
        except ValueError as exc: self.send_error_json(400,str(exc))
        except Exception as exc: self.send_error_json(500,"Internal server error"); print(f"PUT failure: {type(exc).__name__}")


def cookie_name() -> str:
    return "__Host-brandsentinel_session" if os.environ.get("BRANDSENTINEL_ENV")=="production" else COOKIE_NAME


def request_source(handler) -> str | None:
    source=handler.client_address[0] if getattr(handler,"client_address",None) else None
    if source and os.environ.get("BRANDSENTINEL_TRUST_PROXY")=="1":
        candidate=handler.headers.get("X-Real-IP","").strip()
        try: source=str(ipaddress.ip_address(candidate))
        except ValueError: pass
    return source


class BoundedThreadingHTTPServer(ThreadingHTTPServer):
    request_queue_size=64

    def __init__(self,*args,**kwargs):
        self._connection_slots=threading.BoundedSemaphore(64)
        super().__init__(*args,**kwargs)

    def process_request(self,request,client_address):
        if not self._connection_slots.acquire(blocking=False):
            try: request.sendall(b"HTTP/1.1 503 Service Unavailable\r\nConnection: close\r\nContent-Length: 0\r\n\r\n")
            except OSError: pass
            self.shutdown_request(request)
            return
        try: super().process_request(request,client_address)
        except Exception:
            self._connection_slots.release()
            raise

    def process_request_thread(self,request,client_address):
        try: super().process_request_thread(request,client_address)
        finally: self._connection_slots.release()


def main():
    parser=argparse.ArgumentParser(description="BrandSentinel governance platform")
    parser.add_argument("command",nargs="?",choices=["serve","setup-admin"],default="serve")
    args=parser.parse_args()
    get_master_key()
    connect().close()
    if args.command=="setup-admin":
        email=input("Administrator email: ").strip()
        name=input("Administrator name: ").strip()
        workspace=input("Workspace name [My organization]: ").strip() or "My organization"
        password=getpass.getpass("Password (12+ characters): ")
        confirm=getpass.getpass("Confirm password: ")
        if password!=confirm: raise SystemExit("Passwords do not match")
        try: create_initial_admin(email,name,password,workspace)
        except ValueError as exc: raise SystemExit(str(exc))
        print("Initial administrator and workspace created. Start the server with: python server.py serve")
        return
    host=os.environ.get("BRANDSENTINEL_HOST","127.0.0.1")
    port=int(os.environ.get("BRANDSENTINEL_PORT","8765"))
    if os.environ.get("BRANDSENTINEL_ENV")=="production" and host in {"127.0.0.1","localhost"}:
        raise SystemExit("Set BRANDSENTINEL_HOST to the private container interface in production; terminate TLS at a trusted reverse proxy.")
    if os.environ.get("BRANDSENTINEL_ENV")=="production" and not os.environ.get("BRANDSENTINEL_ALLOWED_HOSTS"):
        raise SystemExit("Set BRANDSENTINEL_ALLOWED_HOSTS to the public host names served by the trusted reverse proxy.")
    if os.environ.get("BRANDSENTINEL_ENV")=="production" and os.environ.get("BRANDSENTINEL_TRUST_PROXY")!="1":
        raise SystemExit("Set BRANDSENTINEL_TRUST_PROXY=1 and configure the trusted proxy to overwrite X-Real-IP.")
    server=BoundedThreadingHTTPServer((host,port),SecureHandler)
    server.daemon_threads=True
    print(f"BrandSentinel: http://{host}:{port} (Ctrl+C to stop)")
    try: server.serve_forever()
    except KeyboardInterrupt: pass
    finally: server.server_close()


if __name__ == "__main__": main()
